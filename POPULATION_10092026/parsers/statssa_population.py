"""Stats SA mid-year population estimates (P0302) — provincial series.

Parser for the "Provincial projection by sex and age (2002-2026)_web.xlsx"
workbook: a single sheet, wide by year, with a three-column key
(province, sex, age band) and one column per mid-year 2002..2026.

    Name          | Sex  | Age  | 2002 | 2003 | ... | 2026
    Eastern Cape  | Male | 0-4  | ...  | ...  | ... | ...

We emit one tidy row per (province, sex, age band, year) as a population head
`count`. The 2026 column equals Stats SA's official MYPE 2026 estimate (verified
against the release's "Provincial est by age and sex" table), so the whole
2002-2026 span is the published mid-year estimate series, not a forward forecast.

The sheet also carries several unrelated side-tables lower down (an inter-
provincial migration matrix, dependency ratios) that reuse the same columns; we
keep only rows whose three keys are a real (province, sex, age-band) triple, which
rejects that junk exactly (9 provinces x 2 sexes x 17 bands = 306 series).

Values are kept exactly as published — Stats SA's cohort model emits fractional
persons; we do not round (that would alter the published figure).
"""
from __future__ import annotations
import pandas as pd
from openpyxl import load_workbook

# Stats SA's nine provinces, spelled as the workbook spells them.
_PROVINCES = {
    "Eastern Cape", "Free State", "Gauteng", "KwaZulu-Natal", "Limpopo",
    "Mpumalanga", "Northern Cape", "North West", "Western Cape",
}
# The 17 five-year age bands, in published order (open-ended top band 80+).
_AGE_BANDS = {
    "0-4", "5-9", "10-14", "15-19", "20-24", "25-29", "30-34", "35-39",
    "40-44", "45-49", "50-54", "55-59", "60-64", "65-69", "70-74", "75-79", "80+",
}
_SEX = {"Male": "male", "Female": "female"}

# The key columns sit in the first three columns; the header row is the first row
# that starts with the literal "Name" label (row 5 in the current release).
_HEADER_KEY = "name"


def _find_header(rows: list[tuple]) -> int:
    for i, r in enumerate(rows):
        if r and isinstance(r[0], str) and r[0].strip().lower().startswith(_HEADER_KEY):
            return i
    raise ValueError("statssa_population: header row (starting 'Name') not found")


def parse(local_path: str) -> pd.DataFrame:
    wb = load_workbook(local_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    h = _find_header(rows)
    header = rows[h]
    # year columns are the 4-digit integers/strings from column 3 onward
    year_cols: list[tuple[int, str]] = []
    for c in range(3, len(header)):
        v = header[c]
        if v is None:
            continue
        s = str(v).strip().split(".")[0]
        if s.isdigit() and len(s) == 4:
            year_cols.append((c, s))
    if not year_cols:
        raise ValueError("statssa_population: no year columns found in header")

    out = []
    for r in rows[h + 1:]:
        if not r or r[0] is None:
            continue
        prov = str(r[0]).strip()
        sex = str(r[1]).strip() if r[1] is not None else ""
        age = str(r[2]).strip() if r[2] is not None else ""
        # keep only real (province, sex, age-band) triples — rejects the migration
        # matrix / dependency-ratio side-tables stacked lower in the same sheet
        if prov not in _PROVINCES or sex not in _SEX or age not in _AGE_BANDS:
            continue
        for c, year in year_cols:
            val = r[c]
            if val is None or str(val).strip() == "":
                continue
            out.append({
                "series_type": "estimate",
                "sex": _SEX[sex],
                "age_group": age,
                "geography": prov,
                "period": year,
                "frequency": "annual",
                "measure": "count",
                "value": float(val),
                "unit": "persons",
                "series_code": "P0302",
            })

    df = pd.DataFrame(out)
    if df.empty:
        raise ValueError("statssa_population: no population rows parsed")
    return df
