"""Malawi NSO 2018 census — Series A population tables (clean Excel, Tier-2).

Two tables from the "Series A. Population Tables" workbook:

* A1 — population by sex at national, regional and district level (2018 and 2008);
* A4 — population by age group at national and regional level (2018).

A1 gives the sub-national head counts by sex (all ages); A4 gives the national
and regional age structure. Everything is emitted as published (NSO's totals),
"Malawi" mapped to the "Total country" geography label.
"""
from __future__ import annotations
import pandas as pd
from openpyxl import load_workbook


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _geo(name: str) -> str:
    name = name.strip()
    return "Total country" if name.lower() == "malawi" else name


def _age(label: str) -> str:
    label = label.strip()
    return "<1" if label.lower().startswith("less than 1") else label


def _row(sex, age, geo, period, value):
    return {"series_type": "census", "sex": sex, "age_group": age, "geography": geo,
            "period": period, "frequency": "annual", "measure": "count",
            "value": value, "unit": "persons", "series_code": "MPHC2018"}


def parse(local_path: str) -> pd.DataFrame:
    wb = load_workbook(local_path, read_only=True, data_only=True)
    rows = []

    # --- A1: geography x sex, all ages, for 2018 (cols B/C/D) and 2008 (E/F/G) ---
    for r in wb["A1"].iter_rows(min_row=5, values_only=True):
        if not r or not r[0] or not str(r[0]).strip():
            continue
        geo = _geo(str(r[0]))
        for period, (ct, cm, cf) in (("2018", (1, 2, 3)), ("2008", (4, 5, 6))):
            total, male, female = _num(r[ct]), _num(r[cm]), _num(r[cf])
            if total is None:
                continue
            for sex, val in (("total", total), ("male", male), ("female", female)):
                if val is not None:
                    rows.append(_row(sex, "Total", geo, period, val))

    # --- A4: national + regional age structure, 2018 (sex = total) ---
    a4 = list(wb["A4"].iter_rows(values_only=True))
    # the geography header is the early row with the most non-empty label cells
    hdr = max(a4[:6], key=lambda row: sum(1 for v in row[1:8] if v and str(v).strip()),
              default=None)
    if hdr and sum(1 for v in hdr[1:8] if v and str(v).strip()) >= 2:
        cols = [(i, _geo(str(v).replace("Total", "").strip()))
                for i, v in enumerate(hdr) if i >= 1 and v and str(v).strip()]
        for r in a4:
            if not r or not r[0] or not str(r[0]).strip():
                continue
            age = _age(str(r[0]))
            if age.lower() in ("age group", "table a4"):
                continue
            for i, geo in cols:
                val = _num(r[i]) if i < len(r) else None
                if val is not None:
                    rows.append(_row("total", age, geo, "2018", val))

    df = pd.DataFrame(rows)
    if df.empty:
        raise ValueError("malawi_nso_population: no rows parsed")
    return df.drop_duplicates(["geography", "sex", "age_group", "period"])
