"""Parser for Namibia NSA 2023 census population (NPHC 2023).

Reads two clean count tables from the census main-report workbook:

* Table 2.2 ("Population and percent distribution by sex and area") — the national
  total and each region as Total / Male / Female, at all ages;
* Table 2.5 ("Population Distribution by Single years") — the national population
  by single year of age (0..120) and sex.

Everything is a census head count published by NSA; the all-ages and both-sexes
figures are NSA's own totals, emitted as published (not summed here). We only
normalise the controlled-vocabulary labels: the national area "Namibia" ->
"Total country", and the all-ages row -> age_group "Total". The urban/rural rows
of Table 2.2 are a locality split (deferred), so they are skipped.
"""
from __future__ import annotations
import pandas as pd
from openpyxl import load_workbook

# Table 2.2 rows that are locality totals, not a geography — skipped in v1.
_LOCALITY = {"urban", "rural"}


def _sheet(wb, prefix: str):
    for sn in wb.sheetnames:
        if sn.strip().lower().startswith(prefix.strip().lower()):
            return wb[sn]
    raise KeyError(f"namibia_nsa_population: no sheet starting {prefix!r}")


def _num(v):
    """Coerce a cell to a number, or None. Handles ints, floats and comma/space
    grouped strings; rejects merged-cell junk."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace(" ", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def _row(sex, age_group, geography, value):
    return {
        "series_type": "census",
        "sex": sex,
        "age_group": age_group,
        "geography": geography,
        "period": "2023",
        "frequency": "annual",
        "measure": "count",
        "value": value,
        "unit": "persons",
        "series_code": "NPHC2023",
    }


def parse(local_path: str) -> pd.DataFrame:
    wb = load_workbook(local_path, read_only=True, data_only=True)
    rows: list[dict] = []

    # --- Table 2.2: area x sex, all ages (cols B=area, C=Total, D=Male, E=Female) ---
    for r in _sheet(wb, "Table 2.2").iter_rows(min_row=4, values_only=True):
        area = r[1]
        if not area or not str(area).strip():
            continue
        area = str(area).strip()
        total, male, female = _num(r[2]), _num(r[3]), _num(r[4])
        if total is None:                       # not a data row
            continue
        if area.lower() in _LOCALITY:            # urban/rural = locality, defer
            continue
        geography = "Total country" if area.lower() == "namibia" else area
        for sex, val in (("total", total), ("male", male), ("female", female)):
            if val is not None:
                rows.append(_row(sex, "Total", geography, val))

    # --- Table 2.5: national single-year age x sex
    #     (cols B=Years, C=Male, D=Female, E=Total) ---
    for r in _sheet(wb, "Table 2.5").iter_rows(min_row=4, values_only=True):
        yr = r[1]
        if yr is None:
            continue
        age = _num(yr)
        if age is None or age != int(age):       # skip 'Total'/junk rows
            continue
        age_group = str(int(age))
        male, female, total = _num(r[2]), _num(r[3]), _num(r[4])
        for sex, val in (("total", total), ("male", male), ("female", female)):
            if val is not None:
                rows.append(_row(sex, age_group, "Total country", val))

    if not rows:
        raise ValueError("namibia_nsa_population: no population rows parsed")
    return pd.DataFrame.from_records(rows)
