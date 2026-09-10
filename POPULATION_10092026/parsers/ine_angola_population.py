"""INE Angola Censo 2024 — population by province, sex and broad age group (Excel).

Parser for the national census workbook (Censo2024_Angola_INE_Completo.xlsx):

* "População" sheet — province × (Pop. Total, Homens, Mulheres) for the 21 provinces;
* "Resumo Nacional" sheet — the national totals (Total Habitantes, Mulheres, Homens)
  and the three broad age groups (0-14, 15-64, 65+).

We emit population by province and sex (age_group "Total"), the national total by
sex, and the national broad age groups. Portuguese labels: Homens = male,
Mulheres = female. Everything as published (2024 census).
"""
from __future__ import annotations
import re
import pandas as pd
from openpyxl import load_workbook

_POP = "População"
_RES = "Resumo Nacional"
# age-band labels use a mojibake dash; match by digits (dash-agnostic).
_AGE_BAND = re.compile(r"^(\d{1,2})\D{1,3}(\d{1,2})\s*anos|^(\d{1,2})\+\s*anos", re.I)


def _num(v):
    if v is None:
        return None
    try:
        return float(str(v).replace(",", "").replace("%", "").strip())
    except (TypeError, ValueError):
        return None


def _rows(geo, age, m, f, t):
    return [{"series_type": "census", "sex": sex, "age_group": age, "geography": geo,
             "period": "2024", "frequency": "annual", "measure": "count",
             "value": v, "unit": "persons", "series_code": "AO_CENSO2024"}
            for sex, v in (("total", t), ("male", m), ("female", f)) if v is not None]


def _sheet(wb, name):
    for sn in wb.sheetnames:
        if sn.strip() == name or sn.strip().lower() == name.lower():
            return wb[sn]
    # tolerate mojibake in the sheet name (e.g. "Popula��o")
    key = name.split("ç")[0].lower()
    return next((wb[s] for s in wb.sheetnames if s.lower().startswith(key)), None)


def parse(local_path: str) -> pd.DataFrame:
    wb = load_workbook(local_path, read_only=True, data_only=True)
    out = []

    # --- provinces x sex ---
    ws = _sheet(wb, _POP)
    for r in (ws.iter_rows(values_only=True) if ws else []):
        name = str(r[0]).strip() if r and r[0] is not None else ""
        total, male, female = _num(r[2]) if len(r) > 2 else None, \
            _num(r[3]) if len(r) > 3 else None, _num(r[4]) if len(r) > 4 else None
        if not name or total is None or name.lower().startswith(("prov", "censo")):
            continue
        out += _rows(name, "Total", male, female, total)

    # --- national totals + broad age groups ---
    res = _sheet(wb, _RES)
    nat = {}
    for r in (res.iter_rows(values_only=True) if res else []):
        label = str(r[0]).strip() if r and r[0] is not None else ""
        val = _num(r[1]) if len(r) > 1 else None
        if val is None:
            continue
        low = label.lower()
        if low.startswith("total habitantes"):
            nat["total"] = val
        elif low.startswith("mulheres"):
            nat["female"] = val
        elif low.startswith("homens"):
            nat["male"] = val
        elif "anos" in low:
            m = _AGE_BAND.match(label)
            if m:
                age = f"{m.group(3)}+" if m.group(3) else f"{m.group(1)}-{m.group(2)}"
                out.append({"series_type": "census", "sex": "total", "age_group": age,
                            "geography": "Total country", "period": "2024",
                            "frequency": "annual", "measure": "count", "value": val,
                            "unit": "persons", "series_code": "AO_CENSO2024"})
    if nat.get("total"):
        out += _rows("Total country", "Total", nat.get("male"), nat.get("female"), nat["total"])

    df = pd.DataFrame(out)
    if df.empty:
        raise ValueError("ine_angola_population: no rows parsed")
    return df.drop_duplicates(["geography", "age_group", "sex", "period"])
