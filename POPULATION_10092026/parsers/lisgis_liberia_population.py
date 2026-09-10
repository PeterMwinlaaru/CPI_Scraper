"""LISGIS Liberia 2022 census — population by county and sex (clean Excel).

Parser for the LISGIS county_population.xlsx (2022 Population & Housing Census):
one row per county with Total / Male / Female. We emit those as published (age not
broken down in this file) plus a national "Total country" row summed from the 15
counties (an exhaustive-parts aggregation of the census counts).
"""
from __future__ import annotations
import pandas as pd
from openpyxl import load_workbook


def _num(v):
    if v is None:
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def _rows(geo, m, f, t):
    return [{"series_type": "census", "sex": sex, "age_group": "Total",
             "geography": geo, "period": "2022", "frequency": "annual",
             "measure": "count", "value": v, "unit": "persons",
             "series_code": "LR_PHC2022"} for sex, v in
            (("total", t), ("male", m), ("female", f)) if v is not None]


def parse(local_path: str) -> pd.DataFrame:
    wb = load_workbook(local_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    out, sm, sf, st = [], 0.0, 0.0, 0.0
    for r in rows[1:]:
        name = str(r[0]).strip() if r and r[0] is not None else ""
        total, male, female = _num(r[1]), _num(r[2]), _num(r[3])
        if not name or total is None:
            continue
        out += _rows(name, male, female, total)
        st += total
        sm += male or 0
        sf += female or 0
    if not out:
        raise ValueError("lisgis_liberia_population: no county rows parsed")
    # national total = sum of the (exhaustive) census counties
    out += _rows("Total country", sm, sf, st)
    return pd.DataFrame.from_records(out)
